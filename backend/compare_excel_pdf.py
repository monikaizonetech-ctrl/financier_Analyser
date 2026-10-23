import pdfplumber
import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\IZONE 194\Downloads\GSTR-1-2018-2019.xlsx', data_only=True)
ws = wb.active
print("Title in Excel cell 1,1:", ws.cell(1,1).value)
print("Row 3 in Excel:", ws.cell(3,1).value)

with pdfplumber.open(r'uploads\c84379c7-9ffd-47e8-83b5-1a71b44cb6ea\GSTR-1-2018-2019.pdf') as pdf:
    p0 = pdf.pages[0].extract_text()
    print("Page 0 text lines:")
    for line in p0.split("\n")[:10]:
        print("  ", line)
