"""
Dedicated regression test suite for real-world Karur Vysya Bank (KVB) bank statements.
Verifies dynamic extraction, accurate transaction counts, balance reconciliation,
and guarantees that no mock/demo or cross-statement data leaks into analytics or exports.
"""
import io
import os
import unittest
import openpyxl
import pdfplumber

from app.services.pdf_extractor import extract_pdf_document, extract_document_metadata
from app.services.transaction_parser import parse_document_transactions
from app.services.reconciliation_service import reconcile_transactions, ValidationStatus
from app.services.file_analysis_service import (
    analyse_bank_statement,
    run_analysis,
    build_pdf_report,
    build_excel_report,
    ReportType,
)


class TestKVBBankStatement(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.artifacts_dir = os.path.join(os.path.dirname(__file__), "artifacts")
        cls.ksquare_pdf = os.path.join(cls.artifacts_dir, "sample.pdf")
        cls.janani_pdf = os.path.join(cls.artifacts_dir, "janani_ka_input.pdf")

    def test_01_ksquare_statement_dynamic_extraction(self):
        """Verify dynamic extraction and reconciliation for K Square statement without hardcoding."""
        if not os.path.exists(self.ksquare_pdf):
            self.skipTest(f"sample.pdf not found at {self.ksquare_pdf}")

        doc = extract_pdf_document(self.ksquare_pdf, max_pages=None)
        self.assertTrue(doc.is_valid)

        meta = extract_document_metadata(doc)
        self.assertEqual(meta["bank_name"], "KARUR VYSYA BANK")
        self.assertIn("K SQUARE", meta["account_holder"])
        self.assertEqual(meta["account_no"], "769105000299")

        raw_txns = parse_document_transactions(doc, "sample.pdf", meta)
        self.assertEqual(len(raw_txns), 660)

        # Reconcile
        rec = reconcile_transactions(raw_txns)
        self.assertEqual(rec.status, ValidationStatus.VALID)
        self.assertEqual(len(rec.discrepancies), 0)

        # Run full analysis
        result = run_analysis(ReportType.BSA, [self.ksquare_pdf])
        analytics = result["analytics"]
        self.assertEqual(len(analytics["all_rows"]), 660)
        self.assertEqual(analytics["validation_status"], "VALID")

        # Verify no demo/mock names or wrong bank names leak
        analytics_str = str(analytics)
        self.assertNotIn("DEMO NATIONAL BANK", analytics_str)
        self.assertNotIn("Monika L", analytics_str)
        self.assertNotIn("Mock User", analytics_str)
        self.assertNotIn("Sample Bank Statement", analytics_str)

        # Generate reports
        pdf_bytes = build_pdf_report("KSquare_Report", "BSA", result)
        excel_bytes = build_excel_report("KSquare_Report", "BSA", result)
        self.assertGreater(len(pdf_bytes), 20000)
        self.assertGreater(len(excel_bytes), 20000)

        # Verify PDF contains the actual bank name and account holder
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            p1_text = pdf.pages[0].extract_text()
            self.assertIn("KARUR VYSYA BANK", p1_text)
            self.assertIn("K SQUARE", p1_text)
            self.assertNotIn("DEMO NATIONAL BANK", p1_text)
            self.assertNotIn("Monika L", p1_text)

        # Verify Excel sheet count and Overview content
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertEqual(len(wb.sheetnames), 28)
        ws_ov = wb["Overview"]
        ov_text = str(ws_ov["A5"].value)
        self.assertIn("KARUR VYSYA BANK", ov_text)
        self.assertIn("K SQUARE", ov_text)
        self.assertNotIn("Monika L", ov_text)

    def test_02_janani_statement_dynamic_extraction(self):
        """Verify dynamic extraction and reconciliation for Janani statement."""
        if not os.path.exists(self.janani_pdf):
            self.skipTest(f"janani_ka_input.pdf not found at {self.janani_pdf}")

        doc = extract_pdf_document(self.janani_pdf, max_pages=None)
        self.assertTrue(doc.is_valid)

        meta = extract_document_metadata(doc)
        self.assertEqual(meta["bank_name"], "KARUR VYSYA BANK")
        self.assertIn("JANANI", meta["account_holder"].upper())

        raw_txns = parse_document_transactions(doc, "janani_ka_input.pdf", meta)
        self.assertEqual(len(raw_txns), 381)

        # Verify no uncertain transactions
        uncertain = [t for t in raw_txns if t.get("is_uncertain")]
        self.assertEqual(len(uncertain), 0)

        # Financial verification
        total_cr = round(sum(float(t.get("credit") or 0.0) for t in raw_txns), 2)
        total_dr = round(sum(float(t.get("debit") or 0.0) for t in raw_txns), 2)
        self.assertAlmostEqual(total_cr, 439951.00, places=2)
        self.assertAlmostEqual(total_dr, 412963.51, places=2)

        # Reconcile
        rec = reconcile_transactions(raw_txns)
        self.assertEqual(rec.status, ValidationStatus.VALID)
        self.assertEqual(len(rec.discrepancies), 0)

        # Run full analysis
        result = run_analysis(ReportType.BSA, [self.janani_pdf])
        analytics = result["analytics"]
        self.assertEqual(len(analytics["all_rows"]), 381)
        self.assertEqual(analytics["validation_status"], "VALID")
        self.assertAlmostEqual(analytics["total_credit"], 439951.00, places=2)
        self.assertAlmostEqual(analytics["total_debit"], 412963.51, places=2)

        # Verify no cross-statement pollution
        analytics_str = str(analytics)
        self.assertNotIn("DEMO NATIONAL BANK", analytics_str)
        self.assertNotIn("Monika L", analytics_str)
        self.assertNotIn("K SQUARE CONSTRUCTIONS", analytics_str)
        self.assertNotIn("Sample Bank Statement", analytics_str)

        # Generate reports
        pdf_bytes = build_pdf_report("Janani_Report", "BSA", result)
        excel_bytes = build_excel_report("Janani_Report", "BSA", result)
        self.assertGreater(len(pdf_bytes), 20000)
        self.assertGreater(len(excel_bytes), 20000)

        # Verify PDF
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            p1_text = pdf.pages[0].extract_text()
            self.assertIn("KARUR VYSYA BANK", p1_text)
            self.assertIn("JANANI", p1_text)
            self.assertNotIn("K SQUARE CONSTRUCTIONS", p1_text)
            self.assertNotIn("DEMO NATIONAL BANK", p1_text)

        # Verify Excel
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertEqual(len(wb.sheetnames), 28)
        ws_ov = wb["Overview"]
        ov_text = str(ws_ov["A5"].value)
        self.assertIn("KARUR VYSYA BANK", ov_text)
        self.assertIn("JANANI", ov_text)
        self.assertNotIn("K SQUARE CONSTRUCTIONS", ov_text)


if __name__ == "__main__":
    unittest.main()

