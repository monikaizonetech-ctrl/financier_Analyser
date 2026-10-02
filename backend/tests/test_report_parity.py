"""
Comprehensive test suite for Phase 6: Report Parity & Integration.
Verifies that PDF and Excel reports consume the EXACT same canonical transaction dataset,
guaranteeing full mathematical and transactional consistency.
"""
from datetime import date
import io
import os
import unittest

import openpyxl
import pdfplumber

from app.services.file_analysis_service import (
    ReportType,
    run_analysis,
    build_pdf_report,
    build_excel_report,
    _mask_account_no,
)


class TestReportParity(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.demo_pdf_path = os.path.join(
            os.path.dirname(os.path.dirname(__file__)),
            "uploads",
            "09dd73df-f2e1-4cec-9b00-7babdf6b8a54",
            "demo_bank_statement.pdf",
        )
        if not os.path.exists(cls.demo_pdf_path):
            raise FileNotFoundError(f"Demo PDF not found at {cls.demo_pdf_path}")

        # Run analysis through the single canonical pipeline
        cls.analysis_result = run_analysis(
            ReportType.BSA,
            [cls.demo_pdf_path],
            account_holder="Monika L",
        )
        cls.analytics = cls.analysis_result.get("analytics", {})

        # Generate both reports from the exact same analysis_result
        cls.pdf_bytes = build_pdf_report("Test_Statement_Report", "BSA", cls.analysis_result)
        cls.excel_bytes = build_excel_report("Test_Statement_Report", "BSA", cls.analysis_result)

    def test_01_account_masking(self):
        """Verify account masking preserves the last 4 digits while masking previous digits."""
        self.assertEqual(_mask_account_no("123456789012"), "XXXX XXXX 9012")
        self.assertEqual(_mask_account_no("XXXX XXXX 4821"), "XXXX XXXX 4821")
        self.assertEqual(_mask_account_no("4821"), "XXXX 4821")
        self.assertEqual(_mask_account_no(""), "XXXX XXXX XXXX")
        self.assertEqual(_mask_account_no(None), "XXXX XXXX XXXX")

    def test_02_metrics_parity_pdf_and_excel(self):
        """
        Verify that:
        - Opening Balance
        - Total Income
        - Total Expense
        - Closing Balance
        - Net Cash Flow
        - Transaction Count
        are 100% consistent between the PDF report and the Excel report.
        """
        a = self.analytics
        expected_opening = a["opening_balance"]
        expected_income = a["total_credit"]
        expected_expense = a["total_debit"]
        expected_closing = a["closing_balance"]
        expected_net = a["net_cash_flow"]
        expected_count = len(a["all_rows"])

        # 1. Verify Excel metrics from openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(self.excel_bytes))
        ws_sum = wb["Summary"]
        excel_metrics = {}
        for r in range(3, 10):
            label = ws_sum.cell(row=r, column=1).value
            val = ws_sum.cell(row=r, column=3).value
            if label:
                excel_metrics[label] = val

        self.assertAlmostEqual(float(excel_metrics["Opening Balance"]), expected_opening, places=2)
        self.assertAlmostEqual(float(excel_metrics["Total Inflow Amount"]), expected_income, places=2)
        self.assertAlmostEqual(float(excel_metrics["Total Outflow Amount"]), expected_expense, places=2)
        self.assertAlmostEqual(float(excel_metrics["Closing Balance"]), expected_closing, places=2)

        # Net cash flow in Excel Cashflow Summary
        ws_cf = wb["Cashflow Summary"]
        cf_row = 3  # Total / Jul 2026 row
        cf_net = float(ws_cf.cell(row=cf_row, column=4).value)
        self.assertAlmostEqual(cf_net, expected_net, places=2)

        # Transaction count in Excel All Txns sheet
        ws_txns = wb["All Txns"]
        excel_txn_count = ws_txns.max_row - 2  # Subtract 2 header rows
        self.assertEqual(excel_txn_count, expected_count)

        # 2. Verify PDF metrics from Page 1 text
        with pdfplumber.open(io.BytesIO(self.pdf_bytes)) as pdf:
            p1_text = pdf.pages[0].extract_text()
            self.assertIn(f"Opening Balance Rs. {expected_opening:,.2f}", p1_text)
            self.assertIn(f"Total Income (Credit) Rs. {expected_income:,.2f}", p1_text)
            self.assertIn(f"Total Expense (Debit) Rs. {expected_expense:,.2f}", p1_text)
            self.assertIn(f"Net Cash Flow Rs. {expected_net:,.2f}", p1_text)
            self.assertIn(f"Closing Balance Rs. {expected_closing:,.2f}", p1_text)
            self.assertIn(f"Number of Transactions {expected_count}", p1_text)

    def test_03_pdf_page1_overview_required_fields(self):
        """
        Verify Page 1 contains all required statement overview fields:
        - Bank name
        - Account holder
        - Masked account number
        - Statement period
        - Opening balance
        - Closing balance
        - Total income
        - Total expense
        - Net cash flow
        - Number of transactions
        """
        with pdfplumber.open(io.BytesIO(self.pdf_bytes)) as pdf:
            p1_text = pdf.pages[0].extract_text()

            # Required fields
            self.assertIn("Bank Name DEMO NATIONAL BANK", p1_text)
            self.assertIn("Account Holder Monika L", p1_text)
            self.assertIn("Masked Account Number XXXX XXXX 4821", p1_text)
            self.assertIn("Statement Period 01 Jul 2026 to 31 Jul 2026", p1_text)
            self.assertIn("Opening Balance Rs. 50,000.00", p1_text)
            self.assertIn("Closing Balance Rs. 82,557.00", p1_text)
            self.assertIn("Total Income (Credit) Rs. 48,945.00", p1_text)
            self.assertIn("Total Expense (Debit) Rs. 16,388.00", p1_text)
            self.assertIn("Net Cash Flow Rs. 32,557.00", p1_text)
            self.assertIn("Number of Transactions 13", p1_text)

            # Verification badge
            self.assertIn("FINANCIAL INTEGRITY VERIFIED (STATUS: VALID)", p1_text)

    def test_04_pdf_page2_visualizations(self):
        """
        Verify Page 2 contains the Visual Analytics Dashboard (6 balanced cards):
        - Income vs Expense
        - Balance Trend
        - Expense by Category
        - Top 5 Income Sources
        - Payment Mode Analysis
        - Risk & Unusual Activity
        """
        with pdfplumber.open(io.BytesIO(self.pdf_bytes)) as pdf:
            p2_text = pdf.pages[1].extract_text()

            self.assertIn("Visual Analytics", p2_text)
            self.assertIn("Transaction Visualizations", p2_text)
            self.assertIn("Income vs Expense", p2_text)
            self.assertIn("Balance Trend", p2_text)
            self.assertIn("Expense by Category", p2_text)
            self.assertIn("Top 5 Income Sources", p2_text)
            self.assertIn("Payment Mode Analysis", p2_text)
            self.assertIn("Risk & Unusual Activity", p2_text)

    def test_05_pdf_page3_additional_analysis(self):
        """Verify Page 3 contains Cashflow Summary, Chq/ECS Returns, Recurring Debits, Loan Txns, and Month 1 CounterParty Totals."""
        with pdfplumber.open(io.BytesIO(self.pdf_bytes)) as pdf:
            p3_text = pdf.pages[2].extract_text()
            self.assertIn("Additional Financial Analysis", p3_text)
            self.assertIn("Cashflow Summary", p3_text)
            self.assertIn("Chq/ECS Return Txns", p3_text)
            self.assertIn("Recurring Debits", p3_text)
            self.assertIn("Loan Txns", p3_text)
            self.assertIn("CounterParty Totals", p3_text)

    def test_06_pdf_page4_transaction_summary(self):
        """Verify Page 4 contains Transaction Summary and total demo BSA PDF pages == 4."""
        with pdfplumber.open(io.BytesIO(self.pdf_bytes)) as pdf:
            self.assertEqual(len(pdf.pages), 4, "Demo BSA PDF report must have exactly 4 pages")
            
            p4_text = pdf.pages[3].extract_text()
            self.assertIn("Transaction Summary", p4_text)
            self.assertIn("Date", p4_text)
            self.assertIn("Description", p4_text)
            self.assertIn("Debit / Expense", p4_text)
            self.assertIn("Credit / Income", p4_text)
            self.assertIn("Balance (Rs.)", p4_text)
            self.assertIn("Category", p4_text)

    def test_07_excel_canonical_dataset_integrity(self):
        """Verify Excel report contains 28 sheets and exact matching canonical transactions."""
        wb = openpyxl.load_workbook(io.BytesIO(self.excel_bytes))
        self.assertEqual(len(wb.sheetnames), 28)

        # Verify Overview sheet contains Account Holder and Masked Account Number
        ws_overview = wb["Overview"]
        account_block_val = str(ws_overview["A5"].value)
        self.assertIn("Account Holder        - Monika L", account_block_val)
        self.assertIn("Account Number        - XXXX XXXX 4821", account_block_val)
        self.assertIn("DEMO NATIONAL BANK", account_block_val)

        # Verify All Txns sheet rows match canonical dataset
        ws_all = wb["All Txns"]
        for idx, r in enumerate(self.analytics["all_rows"], start=3):
            self.assertEqual(ws_all.cell(row=idx, column=6).value, r["description"])
            self.assertAlmostEqual(
                float(ws_all.cell(row=idx, column=10).value),
                r["balance"],
                places=2,
            )

    def test_08_end_to_end_upload_to_download_workflow(self):
        """
        Verify the complete workflow:
        Create Report -> Upload PDF -> Analyse Statement -> Download PDF & Excel
        """
        import shutil
        from fastapi import UploadFile
        from app.database import SessionLocal
        from app.models.user import User
        from app.models.report import Report, ReportType, ReportStatus
        from app.routers.reports import upload_report_file, analyse_report, download_report

        db = SessionLocal()
        try:
            user = db.query(User).first()
            self.assertIsNotNone(user, "User must exist in database")

            # 1. Create a fresh test report
            test_report = Report(
                name="E2E_Parity_Verification_Report",
                report_type=ReportType.BSA,
                owner_id=user.id,
                status=ReportStatus.NEED_TO_ANALYSE,
            )
            db.add(test_report)
            db.commit()
            db.refresh(test_report)

            # 2. Upload the demo statement PDF
            with open(self.demo_pdf_path, "rb") as f:
                content = f.read()
            upload_file = UploadFile(
                file=io.BytesIO(content),
                filename="demo_bank_statement.pdf",
            )
            rf = upload_report_file(
                report_id=test_report.id,
                file=upload_file,
                current_user=user,
                db=db,
            )
            self.assertEqual(rf.file_name, "demo_bank_statement.pdf")

            # 3. Trigger Analysis
            analysed = analyse_report(
                report_id=test_report.id,
                current_user=user,
                db=db,
            )
            self.assertEqual(analysed.status, ReportStatus.READY_TO_USE)
            self.assertIsNotNone(analysed.result_summary)

            # 4. Download PDF
            pdf_response = download_report(
                report_id=test_report.id,
                file_format="pdf",
                current_user=user,
                db=db,
            )
            import asyncio
            async def _consume(gen):
                chunks = []
                async for ch in gen:
                    chunks.append(ch)
                return b"".join(chunks)

            downloaded_pdf_bytes = asyncio.run(_consume(pdf_response.body_iterator))
            self.assertGreater(len(downloaded_pdf_bytes), 1000)

            # Verify downloaded PDF has 4 pages
            with pdfplumber.open(io.BytesIO(downloaded_pdf_bytes)) as dl_pdf:
                self.assertEqual(len(dl_pdf.pages), 4)
                p1_text = dl_pdf.pages[0].extract_text()
                self.assertIn("DEMO NATIONAL BANK", p1_text)
                self.assertIn("Monika L", p1_text)
                self.assertIn("Rs. 50,000.00", p1_text)
                self.assertIn("Rs. 82,557.00", p1_text)

            # 5. Download Excel
            xls_response = download_report(
                report_id=test_report.id,
                file_format="xlsx",
                current_user=user,
                db=db,
            )
            downloaded_xls_bytes = asyncio.run(_consume(xls_response.body_iterator))
            self.assertGreater(len(downloaded_xls_bytes), 5000)

            # Verify downloaded Excel has 28 sheets and matching figures
            dl_wb = openpyxl.load_workbook(io.BytesIO(downloaded_xls_bytes))
            self.assertEqual(len(dl_wb.sheetnames), 28)
            ws_sum = dl_wb["Summary"]
            self.assertAlmostEqual(float(ws_sum.cell(row=3, column=3).value), 50000.0, places=2)
            self.assertAlmostEqual(float(ws_sum.cell(row=6, column=3).value), 82557.0, places=2)

            # 6. Cleanup test report
            db.delete(test_report)
            db.commit()
            report_dir = os.path.join(os.path.dirname(self.demo_pdf_path), "..", test_report.id)
            if os.path.exists(report_dir):
                shutil.rmtree(report_dir, ignore_errors=True)
        finally:
            db.close()


if __name__ == "__main__":
    unittest.main()
