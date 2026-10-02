import os
import unittest
from datetime import date
from app.services.pdf_extractor import extract_pdf_document, extract_document_metadata
from app.services.transaction_parser import parse_document_transactions
from app.services.reconciliation_service import reconcile_transactions, ValidationStatus
from app.services.file_analysis_service import (
    analyse_bank_statement,
    _build_analytics,
    build_pdf_report,
    build_excel_report,
)
from app.models.report import ReportType


class TestKSquareStatementValidation(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pdf_path = r"C:\Users\IZONE 181\Downloads\299 apr 26 - seb.pdf"
        cls.file_exists = os.path.exists(cls.pdf_path)

    def test_canonical_extraction_and_reconciliation(self):
        """Verify the exact validation targets specified for the K SQUARE CONSTRUCTIONS statement."""
        if not self.file_exists:
            self.skipTest(f"Test PDF not found at {self.pdf_path}")

        doc = extract_pdf_document(self.pdf_path, max_pages=None)
        self.assertTrue(doc.is_valid)
        self.assertEqual(len(doc.pages), 34)

        meta = extract_document_metadata(doc)
        self.assertEqual(meta["bank_name"], "ICICI BANK")
        self.assertEqual(meta["account_holder"], "K SQUARE CONSTRUCTIONS")
        self.assertEqual(meta["account_no"], "769105000299")

        raw_txns = parse_document_transactions(doc, "299 apr 26 - seb.pdf", meta)

        # 1. Total transaction count: exactly 660 canonical transactions
        self.assertEqual(len(raw_txns), 660)

        # 2. Credits and Debits count
        credits = [t for t in raw_txns if (t.get("credit") or 0.0) > 0]
        debits = [t for t in raw_txns if (t.get("debit") or 0.0) > 0]
        self.assertEqual(len(credits), 145)
        self.assertEqual(len(debits), 515)

        # 3. Financial totals
        total_credit = round(sum(float(t.get("credit") or 0.0) for t in raw_txns), 2)
        total_debit = round(sum(float(t.get("debit") or 0.0) for t in raw_txns), 2)
        net_flow = round(total_credit - total_debit, 2)

        self.assertAlmostEqual(total_credit, 42574432.87, places=2)
        self.assertAlmostEqual(total_debit, 42578131.84, places=2)
        self.assertAlmostEqual(net_flow, -3698.97, places=2)

        # 4. Opening and Closing Balances
        first = raw_txns[0]
        last = raw_txns[-1]
        opening_bal = round(first["balance"] - float(first.get("credit") or 0.0) + float(first.get("debit") or 0.0), 2)
        closing_bal = round(last["balance"], 2)

        self.assertAlmostEqual(opening_bal, -4983631.95, places=2)
        self.assertAlmostEqual(closing_bal, -4987330.92, places=2)

        # 5. Row-by-Row Reconciliation
        rec = reconcile_transactions(raw_txns)
        self.assertEqual(rec.status, ValidationStatus.VALID)
        self.assertEqual(rec.total_transactions, 660)
        self.assertEqual(rec.verified_transactions, 660)
        self.assertEqual(len(rec.discrepancies), 0)

    def test_full_analysis_and_integrity_badge(self):
        """Verify that analyse_bank_statement, _build_analytics, PDF and Excel generation work correctly."""
        if not self.file_exists:
            self.skipTest(f"Test PDF not found at {self.pdf_path}")

        res = analyse_bank_statement(self.pdf_path)
        self.assertEqual(res["transactions_found"], 660)
        self.assertAlmostEqual(res["total_credits"], 42574432.87, places=2)
        self.assertAlmostEqual(res["total_debits"], 42578131.84, places=2)
        self.assertAlmostEqual(res["net_cash_flow"], -3698.97, places=2)

        analytics = _build_analytics([res], "K SQUARE CONSTRUCTIONS")
        self.assertEqual(analytics["validation_status"], "VALID")
        self.assertTrue(analytics["is_integrity_verified"])
        self.assertEqual(len(analytics["validation_reasons"]), 0)
        self.assertEqual(analytics["statement_period"], "01 Apr 2026 to 17 Sep 2026")
        self.assertEqual(len(analytics["balance_mismatches"]), 0)
        self.assertEqual(len(analytics["duplicate_txns"]), 0)

        # Fix 1 & 15: All 6 months present in time series
        self.assertEqual(len(analytics["months"]), 6)
        expected_months = ['Apr 2026', 'May 2026', 'Jun 2026', 'Jul 2026', 'Aug 2026', 'Sep 2026']
        self.assertEqual(analytics["months"], expected_months)

        # Fix 1 & 11: Exact monthly figures and net cashflow
        expected_monthly = {
            'Apr 2026': (8361500.00, 8376200.70, -14700.70),
            'May 2026': (8285000.00, 7324236.00, 960764.00),
            'Jun 2026': (7041000.00, 7998083.50, -957083.50),
            'Jul 2026': (5669000.00, 5655630.80, 13369.20),
            'Aug 2026': (8267162.94, 8285785.98, -18623.04),
            'Sep 2026': (4950769.93, 4938194.86, 12575.07),
        }
        for m, (exp_cr, exp_dr, exp_net) in expected_monthly.items():
            md = analytics["monthly"][m]
            self.assertAlmostEqual(md["credit"], exp_cr, places=2)
            self.assertAlmostEqual(md["debit"], exp_dr, places=2)
            self.assertAlmostEqual(round(md["credit"] - md["debit"], 2), exp_net, places=2)

        # Fix 3: Top 5 Income Sources groups by counterparty/source, not categories
        top_srcs = analytics["top_income_sources"]
        self.assertGreaterEqual(len(top_srcs), 5)
        for s in top_srcs[:5]:
            self.assertNotIn(s["source"], ["Materials & Supplies", "Vendor & Contractor Payments"])

        # Fix 4: Expense categories decoupled from payment modes
        for mode in ("IMPS", "RTGS", "NEFT", "UPI"):
            self.assertNotIn(mode, analytics["cat_totals_debit"])

        # Fix 5: Payment Mode Analysis is independent for credit and debit
        pm_cr = analytics["payment_modes_credit"]
        pm_dr = analytics["payment_modes_debit"]
        self.assertTrue(all(k in pm_cr for k in ("RTGS", "IMPS", "NEFT")))
        self.assertTrue(all(k in pm_dr for k in ("RTGS", "IMPS", "NEFT", "BANK CHARGE")))
        self.assertTrue(all("pct_amount" in v and "pct_count" in v for v in pm_cr.values()))
        self.assertTrue(all("pct_amount" in v and "pct_count" in v for v in pm_dr.values()))

        # Fix 6 & 7: Top 5 Expenses and Income Transactions
        top_exp = analytics["top_5_expenses"]
        top_inc = analytics["top_5_income_txns"]
        self.assertEqual(len(top_exp), 5)
        self.assertEqual(len(top_inc), 5)
        self.assertEqual(top_exp[0]["amount"], 2000000.0)
        self.assertEqual(top_exp[0]["counterparty"], "SHRIHARITRADERS")
        self.assertEqual(top_inc[0]["amount"], 2000000.0)

        # Fix 9: Unusual Transaction Analysis uses dynamic signals & neutral terminology
        self.assertIn("unusual_txns", analytics)
        for u in analytics["unusual_txns"]:
            self.assertNotIn("fraud", u.get("type", "").lower())
            self.assertNotIn("fraud", u.get("reason", "").lower())
            self.assertIn(u.get("type"), ["Large Transaction", "High Activity", "Unusual Transaction"])

        # Fix 10: Transaction activity daily & weekday
        self.assertIn("daily_activity", analytics)
        self.assertIn("weekday_activity", analytics)
        self.assertEqual(len(analytics["weekday_activity"]), 7)

        # Fix 12: Category total validation
        cat_cr_sum = round(sum(v["amount"] for v in analytics["cat_totals_credit"].values()), 2)
        cat_dr_sum = round(sum(v["amount"] for v in analytics["cat_totals_debit"].values()), 2)
        self.assertAlmostEqual(cat_cr_sum, analytics["total_credit"], places=2)
        self.assertAlmostEqual(cat_dr_sum, analytics["total_debit"], places=2)

        summary = {
            "report_type": ReportType.BSA.value,
            "files_analysed": 1,
            "generated_at": "2026-09-25T12:00:00",
            "account_holder": "K SQUARE CONSTRUCTIONS",
            "analytics": analytics,
            "details": [res],
        }

        # Build PDF & Excel
        pdf_bytes = build_pdf_report("K SQUARE CONSTRUCTIONS", "BSA", summary)
        self.assertGreater(len(pdf_bytes), 50000)

        # Fix 2 & 14: PDF text inspection
        import pdfplumber
        import io
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            p1_text = pdf.pages[0].extract_text()
            self.assertIn("FINANCIAL INTEGRITY VERIFIED (STATUS: VALID)", p1_text)
            p2_text = pdf.pages[1].extract_text()
            # Verify no misleading "0% ↓" badge in balance trend
            self.assertNotIn("0% ↓", p2_text)

        excel_bytes = build_excel_report("K SQUARE CONSTRUCTIONS", "BSA", summary)
        self.assertGreater(len(excel_bytes), 50000)

        # Fix 13: Excel sheet inspection
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertEqual(len(wb.sheetnames), 28)
        self.assertIn("Overview", wb.sheetnames)
        self.assertIn("Cashflow Summary", wb.sheetnames)
        self.assertIn("Categories", wb.sheetnames)

    def test_validation_warning_when_reconciliation_fails(self):
        """Verify that when reconciliation fails or discrepancies exist, the report shows VALIDATION WARNING."""
        corrupted_rows = [
            {"date": date(2026, 4, 1), "description": "Txn 1", "credit": 10000.0, "debit": 0.0, "balance": 10000.0, "transaction_id": "T1"},
            {"date": date(2026, 4, 2), "description": "Txn 2", "credit": 0.0, "debit": 5000.0, "balance": 99999.0, "transaction_id": "T2"}, # Corrupted balance
        ]
        rec = reconcile_transactions(corrupted_rows)
        self.assertNotEqual(rec.status, ValidationStatus.VALID)

        # Simulating summary with error
        bad_analytics = {
            "has_data": True,
            "validation_status": rec.status.value,
            "balance_mismatches": [d.to_dict() for d in rec.discrepancies],
            "duplicate_txns": [],
            "all_rows": corrupted_rows,
            "bank_name": "TEST BANK",
            "account_holder": "CORRUPTED TEST",
            "account_no": "1234567890",
            "statement_period": "01 Apr 2026 to 02 Apr 2026",
            "duration_days": 2,
            "account_type": "Current",
            "opening_balance": 0.0,
            "closing_balance": 99999.0,
            "total_credit": 10000.0,
            "total_debit": 5000.0,
            "net_cash_flow": 5000.0,
            "credit_count": 1,
            "debit_count": 1,
            "avg_balance": 50000.0,
            "min_balance": 10000.0,
            "max_balance": 99999.0,
            "monthly": {},
        }
        bad_summary = {
            "report_type": ReportType.BSA.value,
            "files_analysed": 1,
            "generated_at": "2026-09-25T12:00:00",
            "account_holder": "CORRUPTED TEST",
            "analytics": bad_analytics,
            "details": [],
        }
        pdf_bytes = build_pdf_report("CORRUPTED TEST", "BSA", bad_summary)
        import pdfplumber
        import io
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            p1_text = pdf.pages[0].extract_text()
            self.assertIn("VALIDATION WARNING", p1_text)
            self.assertNotIn("FINANCIAL INTEGRITY VERIFIED (STATUS: VALID)", p1_text)


if __name__ == "__main__":
    unittest.main()
